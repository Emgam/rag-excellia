"""
Unified parsing layer, dispatched by file extension.

NOTE ON SUBSTITUTIONS: the README describes IBM Docling as the single
parser for PDF/DOCX/PPTX/XLSX/HTML/images. Docling is a multi-GB, GPU-
capable model download and was not available in this build environment,
so this module uses lighter, dependency-thin libraries per format
(pypdf, python-docx, python-pptx, openpyxl, BeautifulSoup) that cover the
same extensions with plain layout/table extraction (no learned layout
model, no built-in OCR-quality figure detection). To restore Docling
1:1, replace the body of `parse_document()` with a Docling `DocumentConverter`
call and keep the same `ParsedBlock` output contract — nothing downstream
(chunking, metadata, indexing) needs to change.
"""
from __future__ import annotations

import csv
import json
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, List, Optional

import xml.etree.ElementTree as ET

import yaml

from core.exceptions import ParsingError, UnsupportedFileTypeError
from core.models import ContentType
from ingestion.extractors.code_extractor import extract_code_units
from multimodal.ocr.ocr_fallback import ocr_image
from utils.logger import get_logger

logger = get_logger(__name__)


@dataclass
class ParsedBlock:
    text: str
    content_type: ContentType = ContentType.TEXT
    section_path: List[str] = field(default_factory=list)
    page_number: Optional[int] = None
    extra: Dict[str, Any] = field(default_factory=dict)


@dataclass
class ParsedDocument:
    file_path: Path
    title: Optional[str]
    blocks: List[ParsedBlock]


CODE_EXTENSIONS = {".py", ".js", ".ts", ".java", ".go", ".rs", ".cpp", ".c", ".h", ".hpp"}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg"}


def parse_document(path: Path) -> ParsedDocument:
    ext = path.suffix.lower()
    try:
        if ext == ".pdf":
            return _parse_pdf(path)
        if ext == ".docx":
            return _parse_docx(path)
        if ext == ".pptx":
            return _parse_pptx(path)
        if ext == ".xlsx":
            return _parse_xlsx(path)
        if ext == ".csv":
            return _parse_csv(path)
        if ext in (".md", ".markdown"):
            return _parse_markdown(path)
        if ext in (".html", ".htm"):
            return _parse_html(path)
        if ext == ".json":
            return _parse_json(path)
        if ext == ".xml":
            return _parse_xml(path)
        if ext in (".yaml", ".yml"):
            return _parse_yaml(path)
        if ext in (".log", ".ini", ".cfg", ".conf", ".txt"):
            return _parse_plaintext(path)
        if ext in CODE_EXTENSIONS:
            return _parse_code(path)
        if ext in IMAGE_EXTENSIONS:
            return _parse_image(path)
    except UnsupportedFileTypeError:
        raise
    except Exception as exc:  # noqa: BLE001
        raise ParsingError(f"Failed to parse {path}: {exc}") from exc

    raise UnsupportedFileTypeError(f"No parser registered for extension {ext}")


# --- Individual format parsers -----------------------------------------------

def _parse_pdf(path: Path) -> ParsedDocument:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    blocks: List[ParsedBlock] = []
    for i, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            blocks.append(ParsedBlock(text=text, page_number=i))
    title = (reader.metadata.title if reader.metadata else None) or path.stem
    if not blocks:
        logger.warning("No extractable text in %s (likely scanned) — OCR fallback not run on PDFs.", path)
    return ParsedDocument(file_path=path, title=title, blocks=blocks)


def _parse_docx(path: Path) -> ParsedDocument:
    import docx

    d = docx.Document(str(path))
    blocks: List[ParsedBlock] = []
    section_path: List[str] = []

    for p in d.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        style = (p.style.name or "").lower() if p.style else ""
        if style.startswith("heading") or style == "title":
            try:
                level = int(style.replace("heading", "").strip() or 1)
            except ValueError:
                level = 1
            section_path = section_path[: level - 1] + [text]
            continue
        blocks.append(ParsedBlock(text=text, section_path=list(section_path)))

    for t_idx, table in enumerate(d.tables):
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        md_table = _rows_to_markdown(rows)
        if md_table:
            blocks.append(ParsedBlock(
                text=md_table, content_type=ContentType.TABLE,
                section_path=list(section_path), extra={"table_index": t_idx},
            ))

    return ParsedDocument(file_path=path, title=path.stem, blocks=blocks)


def _parse_pptx(path: Path) -> ParsedDocument:
    from pptx import Presentation

    prs = Presentation(str(path))
    blocks: List[ParsedBlock] = []
    for i, slide in enumerate(prs.slides, start=1):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                texts.append(shape.text_frame.text.strip())
            if shape.has_table:
                rows = [[c.text.strip() for c in row.cells] for row in shape.table.rows]
                md_table = _rows_to_markdown(rows)
                if md_table:
                    blocks.append(ParsedBlock(
                        text=md_table, content_type=ContentType.TABLE,
                        page_number=i, section_path=[f"Slide {i}"],
                    ))
        notes = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        combined = "\n".join(texts)
        if notes:
            combined += f"\n\nSpeaker notes: {notes}"
        if combined.strip():
            blocks.append(ParsedBlock(text=combined.strip(), page_number=i, section_path=[f"Slide {i}"]))
    return ParsedDocument(file_path=path, title=path.stem, blocks=blocks)


def _parse_xlsx(path: Path) -> ParsedDocument:
    import openpyxl

    wb = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
    blocks: List[ParsedBlock] = []
    for sheet in wb.worksheets:
        rows = []
        for row in sheet.iter_rows(values_only=True):
            if any(cell is not None for cell in row):
                rows.append(["" if c is None else str(c) for c in row])
        md_table = _rows_to_markdown(rows)
        if md_table:
            blocks.append(ParsedBlock(
                text=md_table, content_type=ContentType.TABLE,
                section_path=[sheet.title],
            ))
    return ParsedDocument(file_path=path, title=path.stem, blocks=blocks)


def _parse_csv(path: Path) -> ParsedDocument:
    with open(path, newline="", encoding="utf-8", errors="replace") as f:
        rows = list(csv.reader(f))
    md_table = _rows_to_markdown(rows)
    blocks = [ParsedBlock(text=md_table, content_type=ContentType.TABLE)] if md_table else []
    return ParsedDocument(file_path=path, title=path.stem, blocks=blocks)


def _parse_markdown(path: Path) -> ParsedDocument:
    text = path.read_text(encoding="utf-8", errors="replace")
    blocks: List[ParsedBlock] = []
    section_path: List[str] = []
    buffer: List[str] = []

    def flush():
        chunk = "\n".join(buffer).strip()
        if chunk:
            blocks.append(ParsedBlock(text=chunk, section_path=list(section_path)))
        buffer.clear()

    for line in text.splitlines():
        if line.strip().startswith("#"):
            flush()
            level = len(line) - len(line.lstrip("#"))
            heading = line.lstrip("#").strip()
            section_path = section_path[: level - 1] + [heading]
        else:
            buffer.append(line)
    flush()
    return ParsedDocument(file_path=path, title=path.stem, blocks=blocks)


def _parse_html(path: Path) -> ParsedDocument:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(path.read_text(encoding="utf-8", errors="replace"), "html.parser")
    for tag in soup(["script", "style"]):
        tag.decompose()
    title = soup.title.string.strip() if soup.title and soup.title.string else path.stem

    blocks: List[ParsedBlock] = []
    section_path: List[str] = []
    for el in soup.find_all(["h1", "h2", "h3", "h4", "p", "li", "table"]):
        if el.name and el.name.startswith("h"):
            level = int(el.name[1])
            heading = el.get_text(strip=True)
            if heading:
                section_path = section_path[: level - 1] + [heading]
            continue
        if el.name == "table":
            rows = [[c.get_text(strip=True) for c in tr.find_all(["td", "th"])] for tr in el.find_all("tr")]
            md_table = _rows_to_markdown(rows)
            if md_table:
                blocks.append(ParsedBlock(text=md_table, content_type=ContentType.TABLE, section_path=list(section_path)))
            continue
        text = el.get_text(strip=True)
        if text:
            blocks.append(ParsedBlock(text=text, section_path=list(section_path)))
    return ParsedDocument(file_path=path, title=title, blocks=blocks)


def _parse_json(path: Path) -> ParsedDocument:
    data = json.loads(path.read_text(encoding="utf-8", errors="replace"))
    pretty = json.dumps(data, indent=2, ensure_ascii=False)
    return ParsedDocument(file_path=path, title=path.stem, blocks=[ParsedBlock(text=pretty, content_type=ContentType.CODE)])


def _parse_xml(path: Path) -> ParsedDocument:
    ET.parse(str(path))  # validate well-formedness
    text = path.read_text(encoding="utf-8", errors="replace")
    return ParsedDocument(file_path=path, title=path.stem, blocks=[ParsedBlock(text=text, content_type=ContentType.CODE)])


def _parse_yaml(path: Path) -> ParsedDocument:
    data = yaml.safe_load(path.read_text(encoding="utf-8", errors="replace"))
    pretty = yaml.dump(data, sort_keys=False, allow_unicode=True) if data is not None else ""
    return ParsedDocument(file_path=path, title=path.stem, blocks=[ParsedBlock(text=pretty, content_type=ContentType.CODE)])


def _parse_plaintext(path: Path) -> ParsedDocument:
    text = path.read_text(encoding="utf-8", errors="replace")
    return ParsedDocument(file_path=path, title=path.stem, blocks=[ParsedBlock(text=text)])


def _parse_code(path: Path) -> ParsedDocument:
    units = extract_code_units(path)
    blocks = [
        ParsedBlock(
            text=u.text, content_type=ContentType.CODE,
            section_path=[u.name] if u.name else [],
            extra={"unit_type": u.unit_type, "start_line": u.start_line, "end_line": u.end_line},
        )
        for u in units
    ]
    return ParsedDocument(file_path=path, title=path.name, blocks=blocks)


def _parse_image(path: Path) -> ParsedDocument:
    caption = ocr_image(path)
    blocks = [ParsedBlock(text=caption, content_type=ContentType.IMAGE)] if caption else []
    return ParsedDocument(file_path=path, title=path.stem, blocks=blocks)


# --- Helpers -------------------------------------------------------------

def _rows_to_markdown(rows: List[List[str]]) -> str:
    rows = [r for r in rows if any(str(c).strip() for c in r)]
    if not rows:
        return ""
    header, *rest = rows
    lines = ["| " + " | ".join(header) + " |", "| " + " | ".join(["---"] * len(header)) + " |"]
    for r in rest:
        r = r + [""] * (len(header) - len(r))
        lines.append("| " + " | ".join(r[: len(header)]) + " |")
    return "\n".join(lines)
